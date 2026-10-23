"""File conversion and batch tools. Pure Python, no Qt — safe to run in worker threads."""
from __future__ import annotations

import html
import io
import os
from pathlib import Path

from PIL import Image, ImageOps
import pillow_heif

pillow_heif.register_heif_opener()

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".heic", ".heif", ".tif", ".tiff",
              ".avif", ".bmp", ".gif", ".ico", ".jfif"}
PDF_EXTS = {".pdf"}
TEXT_EXTS = {".txt", ".md", ".csv", ".log", ".json", ".xml", ".html", ".htm"}
DOCX_EXTS = {".docx"}

IMAGE_TARGETS = ["PNG", "JPG", "WEBP", "HEIC", "TIFF", "AVIF", "BMP", "PDF", "DOCX"]
PDF_TARGETS = ["DOCX", "JPG", "PNG", "TXT"]
TEXT_TARGETS = ["PDF", "DOCX"]
DOCX_TARGETS = ["PDF", "TXT"]

# Format label -> (file extension, Pillow format name)
IMAGE_FORMATS = {
    "PNG": (".png", "PNG"),
    "JPG": (".jpg", "JPEG"),
    "WEBP": (".webp", "WEBP"),
    "HEIC": (".heic", "HEIF"),
    "TIFF": (".tiff", "TIFF"),
    "AVIF": (".avif", "AVIF"),
    "BMP": (".bmp", "BMP"),
}
EXT_TO_LABEL = {".jpg": "JPG", ".jpeg": "JPG", ".jfif": "JPG", ".png": "PNG", ".webp": "WEBP",
                ".heic": "HEIC", ".heif": "HEIC", ".tif": "TIFF", ".tiff": "TIFF",
                ".avif": "AVIF", ".bmp": "BMP", ".pdf": "PDF", ".docx": "DOCX", ".txt": "TXT"}

# Tool ids -> display labels
TOOL_LABELS = {
    "compress": "Compress", "metadata": "Metadata", "edit": "Edit", "annotate": "Annotate",
    "addbg": "Add BG", "crop": "Crop", "redact": "Redact",
    "clean": "Clean Meta", "merge": "Merge PDF",
}


def kind_of(path: str) -> str | None:
    ext = Path(path).suffix.lower()
    if ext in IMAGE_EXTS:
        return "image"
    if ext in PDF_EXTS:
        return "pdf"
    if ext in TEXT_EXTS:
        return "text"
    if ext in DOCX_EXTS:
        return "docx"
    return None


def targets_for(paths: list[str]) -> list[str]:
    """Conversion targets that apply to every dragged file."""
    result = None
    for p in paths:
        kind = kind_of(p)
        if kind == "image":
            own = EXT_TO_LABEL.get(Path(p).suffix.lower())
            t = [x for x in IMAGE_TARGETS if x != own]
        elif kind == "pdf":
            t = list(PDF_TARGETS)
        elif kind == "text":
            t = list(TEXT_TARGETS)
            if Path(p).suffix.lower() != ".txt":
                t.append("TXT")
        elif kind == "docx":
            t = list(DOCX_TARGETS)
        else:
            return []
        result = t if result is None else [x for x in result if x in t]
    return result or []


def tools_for(paths: list[str]) -> list[str]:
    kinds = {kind_of(p) for p in paths}
    if len(paths) == 1:
        if kinds == {"image"}:
            return ["compress", "metadata", "edit", "annotate", "addbg", "crop", "redact"]
        if kinds == {"pdf"}:
            return ["compress", "metadata"]
        return []
    if kinds <= {"image", "pdf"}:
        return ["compress", "clean", "merge"]
    return []


# ---------------------------------------------------------------- helpers

def unique_path(path: Path) -> Path:
    if not path.exists():
        return path
    i = 2
    while True:
        cand = path.with_name(f"{path.stem} ({i}){path.suffix}")
        if not cand.exists():
            return cand
        i += 1


def derived_path(src: str, suffix_label: str = "", ext: str | None = None) -> Path:
    """'Page 001.jpg' + 'Cropped' -> 'Page 001 Cropped.jpg' next to the source."""
    p = Path(src)
    stem = f"{p.stem} {suffix_label}" if suffix_label else p.stem
    return unique_path(p.with_name(stem + (ext or p.suffix)))


def open_image(path: str) -> Image.Image:
    img = Image.open(path)
    img.load()
    try:
        img = ImageOps.exif_transpose(img)
    except Exception:
        pass
    return img


def flatten(img: Image.Image, bg=(255, 255, 255)) -> Image.Image:
    """RGB copy with transparency composited onto a solid colour."""
    if img.mode in ("RGBA", "LA") or (img.mode == "P" and "transparency" in img.info):
        rgba = img.convert("RGBA")
        base = Image.new("RGB", rgba.size, bg)
        base.paste(rgba, mask=rgba.getchannel("A"))
        return base
    return img.convert("RGB")


def save_image(img: Image.Image, out: Path, label: str | None = None, quality: int = 92,
               exif: bytes | None = None) -> Path:
    """Save a PIL image in the format implied by `label` (or out's extension)."""
    label = label or EXT_TO_LABEL.get(out.suffix.lower(), "PNG")
    fmt = IMAGE_FORMATS[label][1]
    kwargs = {}
    if label in ("JPG", "BMP"):
        img = flatten(img)
    elif img.mode not in ("RGB", "RGBA", "L", "LA"):
        img = img.convert("RGBA" if "A" in img.getbands() or "transparency" in img.info else "RGB")
    if label in ("JPG", "WEBP", "AVIF", "HEIC"):
        kwargs["quality"] = quality
    if label == "JPG":
        kwargs["optimize"] = True
    if label == "PNG":
        kwargs["optimize"] = True
    if label == "TIFF":
        kwargs["compression"] = "tiff_lzw"
    if exif and label in ("JPG", "WEBP", "PNG", "TIFF"):
        kwargs["exif"] = exif
    img.save(out, fmt, **kwargs)
    return out


# ---------------------------------------------------------------- conversions

def convert(src: str, target: str) -> list[Path]:
    kind = kind_of(src)
    if kind == "image":
        return _convert_image(src, target)
    if kind == "pdf":
        return _convert_pdf(src, target)
    if kind == "text":
        return _convert_text(src, target)
    if kind == "docx":
        return _convert_docx(src, target)
    raise ValueError(f"Unsupported file: {src}")


def _convert_image(src: str, target: str) -> list[Path]:
    img = open_image(src)
    if getattr(img, "n_frames", 1) > 1:
        img.seek(0)
    if target in IMAGE_FORMATS:
        out = derived_path(src, ext=IMAGE_FORMATS[target][0])
        return [save_image(img, out, target)]
    if target == "PDF":
        out = derived_path(src, ext=".pdf")
        flatten(img).save(out, "PDF", resolution=150.0)
        return [out]
    if target == "DOCX":
        return [_images_to_docx([img], derived_path(src, ext=".docx"))]
    raise ValueError(target)


def _images_to_docx(images: list[Image.Image], out: Path) -> Path:
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(0.75)
    max_w, max_h = 7.0, 9.5
    for i, img in enumerate(images):
        buf = io.BytesIO()
        flatten(img).save(buf, "PNG")
        buf.seek(0)
        w, h = img.size
        if w / h > max_w / max_h:
            doc.add_picture(buf, width=Inches(max_w))
        else:
            doc.add_picture(buf, height=Inches(max_h))
        if i < len(images) - 1:
            doc.add_page_break()
    doc.save(out)
    return out


def _convert_pdf(src: str, target: str) -> list[Path]:
    import pymupdf

    if target == "DOCX":
        from pdf2docx import Converter

        out = derived_path(src, ext=".docx")
        cv = Converter(src)
        try:
            cv.convert(str(out))
        finally:
            cv.close()
        return [out]

    doc = pymupdf.open(src)
    try:
        if target == "TXT":
            out = derived_path(src, ext=".txt")
            text = "\n\n".join(page.get_text() for page in doc)
            out.write_text(text, encoding="utf-8")
            return [out]
        if target in ("JPG", "PNG"):
            outs = []
            ext = IMAGE_FORMATS[target][0]
            multi = doc.page_count > 1
            for i, page in enumerate(doc):
                pix = page.get_pixmap(dpi=200)
                img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples) if pix.n < 4 \
                    else Image.frombytes("RGBA", (pix.width, pix.height), pix.samples)
                out = derived_path(src, f"- Page {i + 1}" if multi else "", ext)
                outs.append(save_image(img, out, target))
            return outs
    finally:
        doc.close()
    raise ValueError(target)


def _read_text(src: str) -> str:
    raw = Path(src).read_bytes()
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def _lines_to_pdf(lines: list[str], out: Path) -> Path:
    import pymupdf

    body = "".join(
        f"<p style='margin:0 0 4px 0'>{html.escape(l) if l.strip() else '&nbsp;'}</p>" for l in lines
    )
    story = pymupdf.Story(html=f"<body style='font-family:sans-serif;font-size:11pt'>{body}</body>")
    writer = pymupdf.DocumentWriter(str(out))
    mediabox = pymupdf.paper_rect("a4")
    where = mediabox + (56, 56, -56, -56)
    more = True
    while more:
        dev = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()
    return out


def _lines_to_docx(lines: list[str], out: Path) -> Path:
    from docx import Document

    doc = Document()
    for l in lines:
        doc.add_paragraph(l)
    doc.save(out)
    return out


def _convert_text(src: str, target: str) -> list[Path]:
    lines = _read_text(src).splitlines()
    if target == "PDF":
        return [_lines_to_pdf(lines, derived_path(src, ext=".pdf"))]
    if target == "DOCX":
        return [_lines_to_docx(lines, derived_path(src, ext=".docx"))]
    if target == "TXT":
        out = derived_path(src, ext=".txt")
        out.write_text("\n".join(lines), encoding="utf-8")
        return [out]
    raise ValueError(target)


def _docx_lines(src: str) -> list[str]:
    from docx import Document

    doc = Document(src)
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append("\t".join(c.text for c in row.cells))
    return lines


def _convert_docx(src: str, target: str) -> list[Path]:
    lines = _docx_lines(src)
    if target == "TXT":
        out = derived_path(src, ext=".txt")
        out.write_text("\n".join(lines), encoding="utf-8")
        return [out]
    if target == "PDF":
        return [_lines_to_pdf(lines, derived_path(src, ext=".pdf"))]
    raise ValueError(target)


# ---------------------------------------------------------------- batch tools

def compress(src: str) -> list[Path]:
    if kind_of(src) == "pdf":
        import pymupdf

        out = derived_path(src, "Compressed")
        doc = pymupdf.open(src)
        try:
            try:
                doc.rewrite_images(dpi_threshold=150, dpi_target=120, quality=70)
            except Exception:
                pass  # older PyMuPDF or no images
            doc.save(out, garbage=4, deflate=True, clean=True)
        finally:
            doc.close()
        return [out]

    img = open_image(src)
    label = EXT_TO_LABEL.get(Path(src).suffix.lower(), "PNG")
    if label not in IMAGE_FORMATS:
        label = "JPG"
    # Cap very large images at 3000px on the long edge.
    if max(img.size) > 3000:
        img.thumbnail((3000, 3000), Image.LANCZOS)
    out = derived_path(src, "Compressed", IMAGE_FORMATS[label][0])
    if label == "PNG":
        q = img.convert("RGBA").quantize(256, method=Image.Quantize.FASTOCTREE)
        q.save(out, "PNG", optimize=True)
    elif label == "BMP":
        out = derived_path(src, "Compressed", ".png")
        img.save(out, "PNG", optimize=True)
    else:
        save_image(img, out, label, quality=70)
    return [out]


def read_metadata(src: str) -> list[tuple[str, str]]:
    rows: list[tuple[str, str]] = []
    if kind_of(src) == "pdf":
        import pymupdf

        doc = pymupdf.open(src)
        rows.append(("Pages", str(doc.page_count)))
        for k, v in (doc.metadata or {}).items():
            if v:
                rows.append((k, str(v)))
        doc.close()
    else:
        from PIL.ExifTags import TAGS, GPSTAGS

        img = Image.open(src)
        rows += [("Format", str(img.format)), ("Size", f"{img.width} × {img.height}"),
                 ("Mode", img.mode)]
        exif = img.getexif()
        for tag, val in exif.items():
            rows.append((TAGS.get(tag, str(tag)), _fmt(val)))
        for ifd_id in (0x8769,):  # Exif sub-IFD
            try:
                for tag, val in exif.get_ifd(ifd_id).items():
                    rows.append((TAGS.get(tag, str(tag)), _fmt(val)))
            except Exception:
                pass
        try:
            for tag, val in exif.get_ifd(0x8825).items():
                rows.append(("GPS " + GPSTAGS.get(tag, str(tag)), _fmt(val)))
        except Exception:
            pass
        for k, v in img.info.items():
            if k not in ("exif", "icc_profile", "xmp") and not k.startswith("jfif") \
                    and isinstance(v, (str, int, float)):
                rows.append((k, _fmt(v)))
    rows.append(("File size", human_size(os.path.getsize(src))))
    return rows


def _fmt(v) -> str:
    if isinstance(v, bytes):
        return f"<{len(v)} bytes>"
    s = str(v)
    return s if len(s) < 200 else s[:200] + "…"


def strip_metadata(src: str) -> list[Path]:
    out = derived_path(src, "Clean")
    if kind_of(src) == "pdf":
        import pymupdf

        doc = pymupdf.open(src)
        doc.set_metadata({})
        try:
            doc.del_xml_metadata()
        except Exception:
            pass
        doc.save(out, garbage=3, deflate=True)
        doc.close()
        return [out]
    clean = open_image(src).copy()
    clean.info = {}  # drops EXIF/XMP/ICC/text chunks; Pillow only writes what we pass
    label = EXT_TO_LABEL.get(Path(src).suffix.lower(), "PNG")
    if label not in IMAGE_FORMATS:
        label, out = "PNG", out.with_suffix(".png")
    save_image(clean, out, label, quality=95)
    return [out]


def merge_pdf(paths: list[str]) -> list[Path]:
    import pymupdf

    out = unique_path(Path(paths[0]).with_name("Merged.pdf"))
    merged = pymupdf.open()
    for p in paths:
        if kind_of(p) == "pdf":
            with pymupdf.open(p) as d:
                merged.insert_pdf(d)
        else:
            buf = io.BytesIO()
            flatten(open_image(p)).save(buf, "PDF", resolution=150.0)
            with pymupdf.open("pdf", buf.getvalue()) as d:
                merged.insert_pdf(d)
    merged.save(out, garbage=3, deflate=True)
    merged.close()
    return [out]


def human_size(n: float) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024
    return f"{n:.1f} GB"
